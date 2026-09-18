# html_to_docx.py
#
# Turns a Zendesk article body (HTML) into a .docx using python-docx.
# Covers what Help Centre articles actually contain: headings, paragraphs,
# bold/italic/underline/code, links, bullet and numbered lists (nested),
# tables, images (as a placeholder line with the URL), horizontal rules.
#
# The title goes in with Word's Title style so the importer's Word parser can
# find it again on the way back in, and the article's own headings keep their
# levels.

import io
from typing import Dict, Optional

from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.shared import Pt, RGBColor

_BLOCK_TAGS = {"p", "div", "h1", "h2", "h3", "h4", "h5", "h6", "ul", "ol", "li",
               "table", "pre", "blockquote", "hr", "br", "img", "section", "article",
               "header", "footer", "figure", "figcaption"}


def html_to_docx_bytes(title: str, html: str, meta: Optional[Dict] = None) -> bytes:
    doc = Document()
    # Word's "Title" style, so the article's own h1..h6 keep their levels.
    # The importer maps Title back to the article title (see zendesk_importer).
    doc.add_paragraph(title or "Untitled", style="Title")

    if meta:
        p = doc.add_paragraph()
        bits = [f"{k}: {v}" for k, v in meta.items() if v not in (None, "", False)]
        run = p.add_run(" | ".join(bits))
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    soup = BeautifulSoup(html or "", "html.parser")
    body = soup.body or soup
    _walk_children(doc, body)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Walking
# ---------------------------------------------------------------------------
def _walk_children(doc, node, list_style: Optional[str] = None, depth: int = 0):
    """Emit children of a block container. Loose inline runs are gathered into a paragraph."""
    pending = []  # inline nodes not yet flushed into a paragraph

    def flush():
        if pending and any(_text_of(n).strip() for n in pending):
            p = doc.add_paragraph()
            for n in pending:
                _add_inline(p, n)
        pending.clear()

    for child in node.children:
        if isinstance(child, NavigableString):
            if str(child).strip():
                pending.append(child)
            continue
        if not isinstance(child, Tag):
            continue
        name = child.name.lower()
        if name in _BLOCK_TAGS:
            flush()
            _emit_block(doc, child, list_style, depth)
        else:
            pending.append(child)
    flush()


def _emit_block(doc, tag: Tag, list_style: Optional[str], depth: int):
    name = tag.name.lower()

    if name in ("h1", "h2", "h3", "h4", "h5", "h6"):
        doc.add_heading(_text_of(tag).strip(), level=int(name[1]))

    elif name in ("p", "div", "section", "article", "header", "footer", "figure", "figcaption"):
        if _has_block_children(tag):
            _walk_children(doc, tag, list_style, depth)
        else:
            text = _text_of(tag)
            if text.strip() or tag.find("img"):
                style = list_style if (list_style and name == "p" and depth) else None
                p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
                _add_inline(p, tag)

    elif name in ("ul", "ol"):
        style = "List Bullet" if name == "ul" else "List Number"
        if depth > 0:
            style += f" {min(depth + 1, 3)}"
        for li in tag.find_all("li", recursive=False):
            _emit_li(doc, li, style, depth)

    elif name == "li":
        _emit_li(doc, tag, list_style or "List Bullet", depth)

    elif name == "table":
        _emit_table(doc, tag)

    elif name == "pre":
        p = doc.add_paragraph()
        run = p.add_run(tag.get_text())
        run.font.name = "Consolas"
        run.font.size = Pt(9)

    elif name == "blockquote":
        for sub in tag.children:
            if isinstance(sub, Tag) and sub.name.lower() in _BLOCK_TAGS:
                _emit_block(doc, sub, list_style, depth)
            elif _text_of(sub).strip():
                p = doc.add_paragraph(style="Intense Quote")
                _add_inline(p, sub)

    elif name == "hr":
        doc.add_paragraph("_" * 40)

    elif name == "img":
        p = doc.add_paragraph()
        _add_inline(p, tag)

    elif name == "br":
        pass


def _emit_li(doc, li: Tag, style: str, depth: int):
    # First the li's own inline content, then any nested lists
    p = doc.add_paragraph(style=style)
    for child in li.children:
        if isinstance(child, Tag) and child.name.lower() in ("ul", "ol"):
            continue
        if isinstance(child, Tag) and child.name.lower() in ("p", "div"):
            _add_inline(p, child)
        else:
            _add_inline(p, child)
    for sub in li.find_all(["ul", "ol"], recursive=False):
        _emit_block(doc, sub, style, depth + 1)


def _emit_table(doc, table: Tag):
    rows = table.find_all("tr")
    if not rows:
        return
    ncols = max(len(r.find_all(["td", "th"], recursive=False)) for r in rows) or 1
    t = doc.add_table(rows=0, cols=ncols)
    t.style = "Table Grid"
    for r in rows:
        cells = r.find_all(["td", "th"], recursive=False)
        row = t.add_row().cells
        for i, c in enumerate(cells[:ncols]):
            para = row[i].paragraphs[0]
            _add_inline(para, c)
            if c.name.lower() == "th":
                for run in para.runs:
                    run.bold = True
    doc.add_paragraph()


# ---------------------------------------------------------------------------
# Inline
# ---------------------------------------------------------------------------
def _add_inline(paragraph, node, fmt: Optional[Dict] = None):
    fmt = dict(fmt or {})
    if isinstance(node, NavigableString):
        text = str(node)
        if text:
            _add_run(paragraph, text, fmt)
        return
    if not isinstance(node, Tag):
        return

    name = node.name.lower()
    if name in ("strong", "b"):
        fmt["bold"] = True
    elif name in ("em", "i"):
        fmt["italic"] = True
    elif name == "u":
        fmt["underline"] = True
    elif name in ("code", "kbd", "samp", "tt"):
        fmt["code"] = True
    elif name in ("s", "strike", "del"):
        fmt["strike"] = True
    elif name == "a":
        href = (node.get("href") or "").strip()
        text = " ".join(_text_of(node).split())
        if href.startswith(("http://", "https://", "mailto:")) and text:
            _add_hyperlink(paragraph, text, href)
        else:
            for child in node.children:
                _add_inline(paragraph, child, fmt)
        return
    elif name == "br":
        paragraph.add_run().add_break()
        return
    elif name == "img":
        src = (node.get("src") or "").strip()
        alt = (node.get("alt") or "").strip()
        _add_run(paragraph, f"[Image{': ' + alt if alt else ''}] {src}", {"italic": True, "size": 8})
        return
    elif name in _BLOCK_TAGS:
        # Block inside an inline context (rare): just dump its text
        for child in node.children:
            _add_inline(paragraph, child, fmt)
        return

    for child in node.children:
        _add_inline(paragraph, child, fmt)


def _add_run(paragraph, text: str, fmt: Dict):
    # Collapse HTML whitespace the way a browser would, but keep one boundary
    # space so "Hello <b>world</b>" does not become "Helloworld"
    if not fmt.get("code"):
        lead = " " if text[:1].isspace() else ""
        trail = " " if text[-1:].isspace() else ""
        text = lead + " ".join(text.split()) + trail
        if not text.strip():
            text = " " if text else ""
        # No leading space at the start of a paragraph
        if not paragraph.runs:
            text = text.lstrip()
    if not text:
        return
    run = paragraph.add_run(text)
    if fmt.get("bold"):
        run.bold = True
    if fmt.get("italic"):
        run.italic = True
    if fmt.get("underline"):
        run.underline = True
    if fmt.get("strike"):
        run.font.strike = True
    if fmt.get("code"):
        run.font.name = "Consolas"
    if fmt.get("link"):
        run.font.color.rgb = RGBColor(0x1F, 0x4E, 0xB4)
        run.underline = True
    if fmt.get("size"):
        run.font.size = Pt(fmt["size"])
    return run


def _add_hyperlink(paragraph, text: str, url: str):
    """A real clickable hyperlink (python-docx has no high-level API for this)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    style = OxmlElement("w:rStyle")
    style.set(qn("w:val"), "Hyperlink")
    rpr.append(style)
    color = OxmlElement("w:color"); color.set(qn("w:val"), "1F4EB4"); rpr.append(color)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
    new_run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def _text_of(node) -> str:
    if isinstance(node, NavigableString):
        return str(node)
    if isinstance(node, Tag):
        return node.get_text(" ")
    return ""


def _has_block_children(tag: Tag) -> bool:
    return any(isinstance(c, Tag) and c.name.lower() in (_BLOCK_TAGS - {"br", "img"}) for c in tag.children)
